from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Emu

root = Path(r"D:\Devpost\SIH26129\docs\demo\sih-ppt-preview\shots")
out = Path(r"D:\Devpost\SIH26129\docs\demo\Samanvay-SIH2026-IDEA.pptx")
desktop = Path(r"D:\SMART INDIA HACKATHON 2026.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for i in range(1, 7):
    png = root / f"render-{i:02d}.png"
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), Emu(0), Emu(0), width=prs.slide_width, height=prs.slide_height)

prs.save(out)
prs.save(desktop)
print("wrote", out)
print("wrote", desktop)
