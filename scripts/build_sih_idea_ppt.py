"""Fill the official SIH 2026 IDEA template (6 slides + drop instructions)."""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "SIH2026-IDEA-Presentation-Format.pptx"
OUT_PPTX = ROOT / "docs" / "demo" / "Samanvay-SIH2026-IDEA.pptx"

NAVY = RGBColor(0x0B, 0x2E, 0x4A)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
SAFFRON = RGBColor(0xC4, 0x56, 0x1A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x3D, 0x4A, 0x54)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF7, 0xF4, 0xEE)
LINE = RGBColor(0xD4, 0xCB, 0xB8)

TEAM = "Samanvay"


def set_runs(shape, lines: list[str], *, size_pt: int = 20, bold: bool = True, color=INK, align=None):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.clear()
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.name = "Arial"
        run.font.color.rgb = color


def hide(shape):
    shape.left = Inches(-12)
    if shape.has_text_frame:
        shape.text_frame.clear()


def add_box(slide, left, top, width, height, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def add_text(slide, left, top, width, height, lines, *, size=13, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Arial"
        run.font.color.rgb = color
    return box


def set_team_ovals(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.name.startswith("Oval"):
                continue
            try:
                shape.fill.solid()
                shape.fill.fore_color.rgb = NAVY
                shape.line.color.rgb = NAVY
            except Exception:
                pass
            try:
                shape.text_frame.word_wrap = False
                shape.text_frame._txBody.bodyPr.set("anchor", "ctr")
            except Exception:
                pass
            set_runs(shape, [TEAM], size_pt=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            for p in shape.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER


def delete_slide(prs, index: int):
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[index]
    rId = sldId.get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)


def fill_title(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "Problem Statement" in shape.text_frame.text:
            set_runs(
                shape,
                [
                    "Problem Statement ID — SIH26129",
                    "Problem Statement Title — System integration and interoperability among",
                    "government digital platforms, resulting in fragmented service delivery",
                    "Theme — Miscellaneous",
                    "Organisation — Government of Maharashtra",
                    "PS Category — Software",
                    "Team ID — (fill from SIH portal after nomination)",
                    "Team Name (Registered on portal) — Samanvay",
                ],
                size_pt=17,
                bold=True,
                color=INK,
            )


def fill_solution(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "Proposed Solution" in shape.text_frame.text:
            hide(shape)
        if shape.has_text_frame and shape.text_frame.text.strip() == "IDEA TITLE":
            set_runs(shape, ["SAMANVAY — FEDERATED INTEROPERABILITY LAYER"], size_pt=22, bold=True, color=NAVY)

    # three thesis cards
    cards = [
        ("Not a new portal", "Sits between existing Maharashtra systems. Does not replace Aaple Sarkar, Mahabhulekh, MahaDBT, Fire e-approval or MPCB/MAITRI."),
        ("Does not store papers", "Certificates stay with the issuer. Samanvay owns links, consent, catalog, pointers, workflow, tracking, hash-chained audit."),
        ("Generic by construction", "Scholarship, licence/NOC and farmer subsidy are three callers. Journey 3 is catalog configuration, not a new Java module."),
    ]
    y = Inches(1.22)
    w = Inches(3.9)
    gap = Inches(0.18)
    x0 = Inches(0.45)
    for i, (title, body) in enumerate(cards):
        x = x0 + i * (w + gap)
        add_box(slide, x, y, w, Inches(1.72), CARD, LINE)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(1.72))
        bar.fill.solid()
        bar.fill.fore_color.rgb = SAFFRON if i == 0 else (TEAL if i == 1 else NAVY)
        bar.line.fill.background()
        add_text(slide, x + Inches(0.22), y + Inches(0.1), w - Inches(0.3), Inches(0.4), [title], size=14, bold=True, color=NAVY)
        add_text(slide, x + Inches(0.22), y + Inches(0.5), w - Inches(0.3), Inches(1.1), [body], size=12, color=MUTED)

    add_text(slide, Inches(0.45), Inches(3.05), Inches(12.4), Inches(0.32), ["How it addresses the problem"], size=13, bold=True, color=SAFFRON)

    bullets = [
        "Citizen facts (income, caste, 7/12, marks, fire NOC) are issued once by the department of record — then reused under consent.",
        "Connectors speak the department’s protocol (REST / SOAP / SFTP / JDBC). DigiLocker is a consent pipe for issued copies, not the source of truth.",
        "One tracking reference across departments; officer desk retries when a source is down; audit proves what was asked, never a stored payload.",
    ]
    add_text(slide, Inches(0.45), Inches(3.35), Inches(12.4), Inches(1.35), ["•  " + b for b in bullets], size=13, color=INK)

    add_text(slide, Inches(0.45), Inches(4.75), Inches(12.4), Inches(0.28), ["Innovation / uniqueness"], size=13, bold=True, color=SAFFRON)
    add_text(
        slide,
        Inches(0.45),
        Inches(5.02),
        Inches(12.4),
        Inches(1.35),
        [
            "•  Control plane decides, data plane moves — no fetch without a 60-second, single-use access grant (typed, not a convention).",
            "•  Federated by default: pointers in the registry, live fetch, provenance on every field (“₹1,85,000 — Tahsildar Haveli”).",
            "•  Machines propose, humans dispose: probabilistic identity matches never auto-link. Build-time Modulith/ArchUnit fail illegal module edges.",
        ],
        size=13,
        color=INK,
    )


def fill_technical(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "Technologies to be used" in shape.text_frame.text:
            hide(shape)

    add_text(slide, Inches(0.45), Inches(1.18), Inches(12.4), Inches(0.28), ["Stack (working prototype on this laptop)"], size=13, bold=True, color=SAFFRON)
    techs = [
        ("Java 21 · Boot 4", "Spring Modulith core"),
        ("Postgres 16", "Flyway · interop state"),
        ("REST SOAP SFTP JDBC", "protocol adapters"),
        ("BPMN 2.0", "Flowable behind a port"),
        ("Keycloak", "OIDC / IdP brokering"),
        ("Resilience4j", "timeout · circuit"),
    ]
    y = Inches(1.48)
    w = Inches(1.95)
    for i, (t, s) in enumerate(techs):
        x = Inches(0.45) + i * Inches(2.1)
        add_box(slide, x, y, w, Inches(0.78), CARD, LINE)
        add_text(slide, x + Inches(0.08), y + Inches(0.08), w - Inches(0.12), Inches(0.38), [t], size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_text(slide, x + Inches(0.08), y + Inches(0.42), w - Inches(0.12), Inches(0.28), [s], size=10, color=MUTED, align=PP_ALIGN.CENTER)

    add_text(slide, Inches(0.45), Inches(2.38), Inches(12.4), Inches(0.28), ["Methodology — departments issue; Samanvay orchestrates the pull"], size=13, bold=True, color=SAFFRON)

    # flow row
    nodes = [
        ("Citizen portals", "Scholarship / Licence / Farmer"),
        ("Samanvay control", "identity · consent · catalog · registry"),
        ("Samanvay data", "connector · orchestration · tracking"),
        ("Issuers (SoR)", "Aaple Sarkar · Bhulekh · Fire · MPCB · MahaDBT"),
    ]
    y = Inches(2.7)
    w = Inches(2.85)
    for i, (t, s) in enumerate(nodes):
        x = Inches(0.4) + i * Inches(3.2)
        add_box(slide, x, y, w, Inches(1.05), NAVY if i in (1, 2) else TEAL)
        add_text(slide, x + Inches(0.1), y + Inches(0.12), w - Inches(0.2), Inches(0.36), [t], size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(slide, x + Inches(0.1), y + Inches(0.5), w - Inches(0.2), Inches(0.48), [s], size=10, color=WHITE, align=PP_ALIGN.CENTER)
        if i < 3:
            arr = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                x + w + Inches(0.04),
                y + Inches(0.38),
                Inches(0.28),
                Inches(0.28),
            )
            arr.fill.solid()
            arr.fill.fore_color.rgb = SAFFRON
            arr.line.fill.background()

    add_text(
        slide,
        Inches(0.45),
        Inches(3.9),
        Inches(12.4),
        Inches(2.4),
        [
            "Process: consent + identity link  →  signed 60s grant  →  adapter fetch (live, not stored)  →  canonical map + provenance  →  track / officer / audit.",
            "DigiLocker (production): MeitY partner API after Aadhaar consent lists issued URIs. Departments push copies; locker is not a login into Fire that then dumps 7/12.",
            "Demo now: labelled DigiLocker sandbox + mock issuer adapters. Offline compose. P5 sequence — Journey 1 concrete, Journey 2 generalizes, Journey 3 is config.",
            "Working prototype: three portal skins, officer retry, issued-record papers, tamper-evident audit explorer.",
        ],
        size=13,
        color=INK,
    )


def fill_feasibility(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "Analysis of the feasibility" in shape.text_frame.text:
            hide(shape)

    cols = [
        (
            NAVY,
            "Feasible now",
            [
                "Core already runs: Java 21 / Boot 4 / Postgres / three journeys.",
                "Heterogeneous mocks are the proof (SOAP, REST, JDBC, SFTP path).",
                "PS constraint met: no replacement of existing department systems.",
                "Onboarding a journey is catalog + connectors, not a rewrite.",
                "Demo is fully offline — venue network is not a dependency.",
            ],
        ),
        (
            SAFFRON,
            "Challenges / risks",
            [
                "Live DigiLocker / Aadhaar needs a MeitY partner account (not on this laptop).",
                "Real department APIs need MoU / NIC / department credentials.",
                "Keycloak brokering and BPMN depth are fiddly in production.",
                "Probabilistic identity merge is unsafe if auto-applied.",
                "Officer queues and SLA if a source stays down.",
            ],
        ),
        (
            TEAL,
            "How we overcome them",
            [
                "Sandbox labelled as sandbox; issuer mocks stand in for SoR APIs.",
                "WorkflowEngine port: BPMN today, YAML DAG fallback in ~2 days.",
                "CHECK + service + RBAC: probabilistic match never becomes an active link.",
                "Degraded mode + officer retry (Revenue unavailable in demo).",
                "Recorded backup walkthrough if live demo fails.",
            ],
        ),
    ]
    y = Inches(1.28)
    w = Inches(3.95)
    for i, (color, title, bullets) in enumerate(cols):
        x = Inches(0.4) + i * Inches(4.2)
        add_box(slide, x, y, w, Inches(5.15), CARD, LINE)
        head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.48))
        head.fill.solid()
        head.fill.fore_color.rgb = color
        head.line.fill.background()
        add_text(slide, x + Inches(0.12), y + Inches(0.08), w - Inches(0.2), Inches(0.36), [title], size=15, bold=True, color=WHITE)
        add_text(slide, x + Inches(0.16), y + Inches(0.62), w - Inches(0.28), Inches(4.6), ["•  " + b for b in bullets], size=12, color=INK)


def fill_impact(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "Potential impact" in shape.text_frame.text:
            hide(shape)

    rows = [
        ("Citizens", "One issued income / caste / 7/12 / marksheet reused across services. One reference number. Consent that can be revoked — next fetch denied; nothing to purge."),
        ("Officers", "One file, not three portals. See which issuer answered. Retry a down department without restarting the citizen journey."),
        ("Departments", "Keep their system of record. Plug in via adapter + mapping. A new scheme is configuration (farmer subsidy already demonstrates this)."),
        ("State / economic", "Avoid a fourth silo or a central data lake. Interoperability layer is cheaper to extend than N×N bilateral integrations."),
        ("Social / privacy", "Sensitive discovery (e.g. caste pointer) is consent-gated. Provenance on every field. Tamper-evident audit for RTI / inquiry."),
        ("Environment", "Fewer duplicate physical certificates and counter visits; digital reuse of already-issued records."),
    ]
    y0 = Inches(1.22)
    for i, (title, body) in enumerate(rows):
        col = i % 2
        row = i // 2
        x = Inches(0.4) + col * Inches(6.4)
        y = y0 + row * Inches(1.72)
        add_box(slide, x, y, Inches(6.15), Inches(1.58), CARD, LINE)
        add_text(slide, x + Inches(0.2), y + Inches(0.12), Inches(5.75), Inches(0.36), [title], size=14, bold=True, color=NAVY)
        add_text(slide, x + Inches(0.2), y + Inches(0.5), Inches(5.75), Inches(0.95), [body], size=12, color=MUTED)


def fill_refs(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "Details / Links" in shape.text_frame.text:
            hide(shape)

    left = [
        "Problem & policy",
        "•  SIH26129 — GoM / MSINS: interoperability without replacing existing systems",
        "•  India Stack / DEPA-style consent artefacts (purpose, requester, revoke)",
        "•  MeitY DigiLocker issued-document model (issuer push + citizen pull)",
        "",
        "Maharashtra systems of record (not scraped; adapted)",
        "•  Aaple Sarkar — aaplesarkar.mahaonline.gov.in — income / caste",
        "•  Mahabhulekh — bhulekh.mahabhumi.gov.in — 7/12 Record of Rights",
        "•  MahaDBT — mahadbt.maharashtra.gov.in — scholarship / farmer DBT",
        "•  Fire e-approval — mahafireservice.gov.in/e-fire.php",
        "•  MAITRI / MPCB — maitri.maharashtra.gov.in — pollution consent",
        "•  MahaVastu / BPMS — mahavastu.maharashtra.gov.in",
    ]
    right = [
        "Standards we implement against",
        "•  BPMN 2.0 for cross-department journeys",
        "•  OIDC / SAML via Keycloak identity brokering",
        "•  REST, SOAP/XML, SFTP/CSV, JDBC (legacy without an API)",
        "•  JSON Schema canonical model + JSONPath mapping DSL",
        "•  Hash-chained audit (tamper-evident ledger)",
        "",
        "This prototype (design + code)",
        "•  docs/architecture/HLD.md and LLD.md in the Samanvay repo",
        "•  docs/demo/ONE_PAGER.md · JUDGE_SCRIPT.md",
        "•  github.com/himanshu2394i/samanvay",
        "",
        "Honest demo boundary: issuer APIs and DigiLocker partner credentials",
        "are mocked / sandboxed until departments and MeitY onboard.",
    ]
    add_box(slide, Inches(0.4), Inches(1.22), Inches(6.15), Inches(5.45), CARD, LINE)
    add_box(slide, Inches(6.7), Inches(1.22), Inches(6.15), Inches(5.45), CARD, LINE)
    add_text(slide, Inches(0.58), Inches(1.38), Inches(5.8), Inches(5.15), left, size=13, color=INK)
    add_text(slide, Inches(6.88), Inches(1.38), Inches(5.8), Inches(5.15), right, size=13, color=INK)


def main():
    if not TEMPLATE.exists():
        raise SystemExit(f"Template missing: {TEMPLATE}")
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, OUT_PPTX)
    prs = Presentation(str(OUT_PPTX))
    fill_title(prs.slides[0])
    fill_solution(prs.slides[1])
    fill_technical(prs.slides[2])
    fill_feasibility(prs.slides[3])
    fill_impact(prs.slides[4])
    fill_refs(prs.slides[5])
    set_team_ovals(prs)
    delete_slide(prs, 6)
    prs.save(str(OUT_PPTX))
    print("Wrote", OUT_PPTX, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
